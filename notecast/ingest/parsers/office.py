"""Parsers for Microsoft Office formats: PowerPoint (.pptx) and Word (.docx).

PptxParser emits one Section per slide (title, body text in reading order,
tables, and speaker notes appended). DocxParser emits one Section per
heading-delimited block, tracking the ancestor heading path.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from pptx import Presentation
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError
from pptx.shapes.group import GroupShape

from notecast.ingest.parsers import register
from notecast.ingest.textutils import normalize_whitespace
from notecast.models import Location, ParsedDocument, Section, SourceType

_GENERIC_PPTX_TITLES = {"powerpoint presentation", ""}

_HEADING_STYLES = {
    "title",
    "heading 1",
    "heading 2",
    "heading 3",
    "heading 4",
    "heading 5",
    "heading 6",
}

_LIST_STYLES = {"list bullet", "list number", "list bullet 2", "list number 2"}


def _heading_level(style_name: str) -> int | None:
    """Return the heading level (0 for Title, 1-6 for Heading N) for a
    paragraph style name, or None if it isn't a heading style.
    """
    name = style_name.strip().lower()
    if name == "title":
        return 0
    if name.startswith("heading "):
        rest = name[len("heading ") :].strip()
        if rest.isdigit():
            return int(rest)
    return None


# ---------------------------------------------------------------------------
# PPTX
# ---------------------------------------------------------------------------


def _clean_slide_text(text: str) -> str:
    """Trim trailing whitespace per line and collapse runs of blank lines,
    without touching leading indentation (which encodes bullet level).
    """
    lines = [line.rstrip() for line in text.replace("\r\n", "\n").split("\n")]
    out = "\n".join(lines).strip()
    return re.sub(r"\n{3,}", "\n\n", out)


def _shape_sort_key(shape: Any) -> tuple[float, float]:
    top = shape.top if shape.top is not None else 0
    left = shape.left if shape.left is not None else 0
    return (top, left)


def _is_title_shape(shape: Any) -> bool:
    try:
        return bool(shape.is_placeholder and shape.placeholder_format.type is not None) and (
            shape.placeholder_format.idx == 0
        )
    except (AttributeError, ValueError):
        return False


def _text_frame_lines(text_frame: Any) -> list[str]:
    lines: list[str] = []
    for paragraph in text_frame.paragraphs:
        text = "".join(run.text for run in paragraph.runs).strip()
        if not text and paragraph.text.strip():
            text = paragraph.text.strip()
        if not text:
            continue
        level = paragraph.level or 0
        indent = "  " * level
        lines.append(f"{indent}- {text}")
    return lines


def _table_lines(table: Any) -> list[str]:
    lines = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        if any(cells):
            lines.append(" | ".join(cells))
    return lines


def _collect_shape_lines(shape: Any, title_shape_id: int | None) -> list[str]:
    """Collect text lines from a shape (recursing into groups), skipping the
    title shape (handled separately, identified by its shape_id).
    """
    if title_shape_id is not None and getattr(shape, "shape_id", None) == title_shape_id:
        return []
    lines: list[str] = []
    if isinstance(shape, GroupShape) or shape.shape_type == 6:  # 6 = MSO_SHAPE_TYPE.GROUP
        sub_shapes = sorted(shape.shapes, key=_shape_sort_key)
        for sub in sub_shapes:
            lines.extend(_collect_shape_lines(sub, title_shape_id))
        return lines
    if getattr(shape, "has_table", False) and shape.has_table:
        lines.extend(_table_lines(shape.table))
        return lines
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        lines.extend(_text_frame_lines(shape.text_frame))
    return lines


def _slide_title(slide: Any) -> tuple[str | None, int | None]:
    """Return (title text, title shape id) for a slide, if it has a title
    placeholder with text.
    """
    title_shape = None
    try:
        title_shape = slide.shapes.title
    except AttributeError:
        title_shape = None
    title_shape_id = getattr(title_shape, "shape_id", None)
    if title_shape is not None and getattr(title_shape, "has_text_frame", False):
        text = title_shape.text_frame.text.strip()
        if text:
            return text, title_shape_id
    return None, title_shape_id


def _slide_notes(slide: Any) -> str:
    if not slide.has_notes_slide:
        return ""
    notes_slide = slide.notes_slide
    if notes_slide is None or notes_slide.notes_text_frame is None:
        return ""
    return notes_slide.notes_text_frame.text.strip()


class PptxParser:
    """Parses .pptx presentations into one Section per slide."""

    suffixes: tuple[str, ...] = (".pptx",)

    def parse(self, path: Path, source_path: str) -> ParsedDocument:
        try:
            presentation = Presentation(str(path))
        except (PptxPackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
            raise ValueError(
                f"Could not open {path.name} as a PowerPoint file (it may be corrupt "
                "or not a valid .pptx file)."
            ) from exc

        sections: list[Section] = []
        first_slide_title: str | None = None
        has_notes = False

        for index, slide in enumerate(presentation.slides, start=1):
            title_text, title_shape_id = _slide_title(slide)
            if title_text and first_slide_title is None:
                first_slide_title = title_text

            body_shapes = sorted(slide.shapes, key=_shape_sort_key)
            body_lines: list[str] = []
            for shape in body_shapes:
                body_lines.extend(_collect_shape_lines(shape, title_shape_id))

            notes_text = _slide_notes(slide)
            if notes_text:
                has_notes = True

            parts: list[str] = []
            if title_text:
                parts.append(title_text)
            if body_lines:
                parts.append("\n".join(body_lines))
            text = "\n\n".join(parts)
            if notes_text:
                text = (
                    f"{text}\n\nSpeaker notes: {notes_text}"
                    if text
                    else (f"Speaker notes: {notes_text}")
                )

            text = _clean_slide_text(text)
            if not text:
                continue

            sections.append(
                Section(
                    text=text,
                    location=Location(slide=index),
                    title=title_text,
                )
            )

        core_title = (presentation.core_properties.title or "").strip()
        if core_title and core_title.lower() not in _GENERIC_PPTX_TITLES:
            title = core_title
        elif first_slide_title:
            title = first_slide_title
        else:
            title = path.stem

        return ParsedDocument(
            source_path=source_path,
            source_type=SourceType.PPTX,
            title=title,
            sections=sections,
            metadata={
                "slide_count": len(presentation.slides),
                "has_notes": has_notes,
            },
        )


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def _paragraph_style_name(paragraph: Paragraph) -> str:
    try:
        return paragraph.style.name or ""
    except AttributeError:
        return ""


def _is_list_paragraph(paragraph: Paragraph) -> bool:
    style_name = _paragraph_style_name(paragraph).strip().lower()
    if style_name in _LIST_STYLES:
        return True
    num_pr = paragraph._p.find(qn("w:pPr") + "/" + qn("w:numPr"))
    return num_pr is not None


def _paragraph_text(paragraph: Paragraph) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""
    if _is_list_paragraph(paragraph):
        return f"- {text}"
    return text


def _cell_text_docx(cell: Any) -> str:
    """A table cell's own paragraph text, plus any nested tables flattened
    into the same string (rows joined with " ; ") so nothing is lost.
    """
    paragraph_text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
    nested_parts = [paragraph_text] if paragraph_text else []
    for nested_table in cell.tables:
        nested_rows = _table_lines_docx(nested_table)
        if nested_rows:
            nested_parts.append(" ; ".join(nested_rows))
    return " ".join(nested_parts).strip()


def _table_lines_docx(table: Table) -> list[str]:
    lines = []
    for row in table.rows:
        cells = [_cell_text_docx(cell) for cell in row.cells]
        if any(cells):
            lines.append(" | ".join(cells))
    return lines


def _iter_body_blocks(document: DocumentObject) -> list[Any]:
    """Iterate the document body's direct XML children, yielding Paragraph
    or Table objects in document order (so tables interleave correctly).
    """
    blocks: list[Any] = []
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            blocks.append(Paragraph(child, document))
        elif child.tag == qn("w:tbl"):
            blocks.append(Table(child, document))
    return blocks


class DocxParser:
    """Parses .docx documents into sections split by heading."""

    suffixes: tuple[str, ...] = (".docx",)

    def parse(self, path: Path, source_path: str) -> ParsedDocument:
        try:
            document = Document(str(path))
        except (zipfile.BadZipFile, KeyError, ValueError, DocxPackageNotFoundError) as exc:
            raise ValueError(
                f"Could not open {path.name} as a Word document (it may be corrupt "
                "or not a valid .docx file)."
            ) from exc

        title_style_text: str | None = None
        first_heading1_text: str | None = None

        heading_stack: list[tuple[int, str]] = []  # (level, text)
        current_lines: list[str] = []
        raw_sections: list[tuple[list[str], list[str]]] = []  # (heading_path, lines)

        def flush() -> None:
            heading_path = [text for _level, text in heading_stack]
            if current_lines:
                raw_sections.append((heading_path.copy(), current_lines.copy()))
            current_lines.clear()

        for block in _iter_body_blocks(document):
            if isinstance(block, Table):
                lines = _table_lines_docx(block)
                current_lines.extend(lines)
                continue

            paragraph = block
            style_name = _paragraph_style_name(paragraph)
            level = _heading_level(style_name)
            text = paragraph.text.strip()

            if level is not None:
                if not text:
                    continue
                flush()
                if level == 0:
                    if title_style_text is None:
                        title_style_text = text
                    # Title-style paragraphs start a fresh top-level section
                    # and don't nest under prior headings.
                    heading_stack = [(0, text)]
                else:
                    if level == 1 and first_heading1_text is None:
                        first_heading1_text = text
                    while heading_stack and heading_stack[-1][0] >= level:
                        heading_stack.pop()
                    heading_stack.append((level, text))
                continue

            line = _paragraph_text(paragraph)
            if line:
                current_lines.append(line)

        flush()

        sections: list[Section] = []
        for heading_path, lines in raw_sections:
            text = normalize_whitespace("\n".join(lines))
            if not text:
                continue
            sections.append(
                Section(
                    text=text,
                    location=Location(heading_path=heading_path),
                    title=heading_path[-1] if heading_path else None,
                )
            )

        core_title = (document.core_properties.title or "").strip()
        if core_title:
            title = core_title
        elif title_style_text:
            title = title_style_text
        elif first_heading1_text:
            title = first_heading1_text
        else:
            title = path.stem

        return ParsedDocument(
            source_path=source_path,
            source_type=SourceType.DOCX,
            title=title,
            sections=sections,
            metadata={},
        )


register(PptxParser())
register(DocxParser())
