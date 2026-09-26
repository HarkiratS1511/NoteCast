"""Tests for the PowerPoint and Word parsers, using synthetic fixtures built
with python-pptx and python-docx (never real course material).
"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from notecast.ingest.parsers import get_parser
from notecast.ingest.parsers.office import DocxParser, PptxParser
from notecast.models import SourceType

# ---------------------------------------------------------------------------
# PPTX fixtures
# ---------------------------------------------------------------------------


def _make_pptx_with_bullets(tmp_path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Introduction to Graphs"
    body = slide.placeholders[1].text_frame
    body.text = "Top level point"
    p2 = body.add_paragraph()
    p2.text = "Nested point"
    p2.level = 1
    p3 = body.add_paragraph()
    p3.text = "Another top level"
    p3.level = 0

    path = tmp_path / "bullets.pptx"
    prs.save(str(path))
    return path


def _make_pptx_with_group_and_table(tmp_path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # title only layout
    slide.shapes.title.text = "Shapes and Tables"

    box1 = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(2), Inches(1))
    box1.text_frame.text = "Grouped shape A"
    box2 = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(1))
    box2.text_frame.text = "Grouped shape B"
    group_shapes = slide.shapes.add_group_shape([box1, box2])
    assert group_shapes is not None

    table_shape = slide.shapes.add_table(2, 2, Inches(4), Inches(2), Inches(3), Inches(1))
    table = table_shape.table
    table.cell(0, 0).text = "Header 1"
    table.cell(0, 1).text = "Header 2"
    table.cell(1, 0).text = "a"
    table.cell(1, 1).text = "b"

    path = tmp_path / "group_table.pptx"
    prs.save(str(path))
    return path


def _make_pptx_with_notes(tmp_path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Notes Slide"
    slide.placeholders[1].text_frame.text = "Body text"
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = "The lecturer says this is important."

    path = tmp_path / "notes.pptx"
    prs.save(str(path))
    return path


def _make_pptx_with_empty_slide(tmp_path: Path) -> Path:
    prs = Presentation()
    slide1 = prs.slides.add_slide(prs.slide_layouts[1])
    slide1.shapes.title.text = "First"
    slide1.placeholders[1].text_frame.text = "Content"

    # Blank layout, no text anywhere, no notes.
    prs.slides.add_slide(prs.slide_layouts[6])

    slide3 = prs.slides.add_slide(prs.slide_layouts[1])
    slide3.shapes.title.text = "Third"
    slide3.placeholders[1].text_frame.text = "More content"

    path = tmp_path / "with_empty.pptx"
    prs.save(str(path))
    return path


def _make_pptx_no_title(tmp_path: Path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank, no placeholders
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    box.text_frame.text = "Just some body text, no title"

    path = tmp_path / "no_title.pptx"
    prs.save(str(path))
    return path


class TestPptxParser:
    def test_title_and_bullet_levels(self, tmp_path: Path) -> None:
        path = _make_pptx_with_bullets(tmp_path)
        doc = PptxParser().parse(path, "week-01/bullets.pptx")

        assert doc.source_type == SourceType.PPTX
        assert len(doc.sections) == 1
        section = doc.sections[0]
        assert section.title == "Introduction to Graphs"
        assert section.location.slide == 1
        assert "Introduction to Graphs" in section.text
        assert "- Top level point" in section.text
        assert "  - Nested point" in section.text
        assert "- Another top level" in section.text

    def test_grouped_shapes_and_table(self, tmp_path: Path) -> None:
        path = _make_pptx_with_group_and_table(tmp_path)
        doc = PptxParser().parse(path, "shapes.pptx")

        assert len(doc.sections) == 1
        text = doc.sections[0].text
        assert "Grouped shape A" in text
        assert "Grouped shape B" in text
        assert "Header 1 | Header 2" in text
        assert "a | b" in text

    def test_speaker_notes_appended(self, tmp_path: Path) -> None:
        path = _make_pptx_with_notes(tmp_path)
        doc = PptxParser().parse(path, "notes.pptx")

        assert doc.metadata["has_notes"] is True
        text = doc.sections[0].text
        assert "Speaker notes: The lecturer says this is important." in text
        assert text.index("Body text") < text.index("Speaker notes:")

    def test_empty_slide_skipped_and_numbering_preserved(self, tmp_path: Path) -> None:
        path = _make_pptx_with_empty_slide(tmp_path)
        doc = PptxParser().parse(path, "with_empty.pptx")

        assert doc.metadata["slide_count"] == 3
        assert len(doc.sections) == 2
        assert [s.location.slide for s in doc.sections] == [1, 3]
        assert doc.sections[0].title == "First"
        assert doc.sections[1].title == "Third"

    def test_title_fallback_to_first_slide_title(self, tmp_path: Path) -> None:
        path = _make_pptx_with_bullets(tmp_path)
        doc = PptxParser().parse(path, "bullets.pptx")
        assert doc.title == "Introduction to Graphs"

    def test_title_fallback_to_stem_when_no_title_anywhere(self, tmp_path: Path) -> None:
        path = _make_pptx_no_title(tmp_path)
        doc = PptxParser().parse(path, "no_title.pptx")
        assert doc.title == "no_title"

    def test_core_properties_title_used_when_present(self, tmp_path: Path) -> None:
        path = _make_pptx_with_bullets(tmp_path)
        prs = Presentation(str(path))
        prs.core_properties.title = "COMP4650 Week 1"
        prs.save(str(path))

        doc = PptxParser().parse(path, "bullets.pptx")
        assert doc.title == "COMP4650 Week 1"

    def test_generic_core_title_falls_back(self, tmp_path: Path) -> None:
        path = _make_pptx_with_bullets(tmp_path)
        prs = Presentation(str(path))
        prs.core_properties.title = "PowerPoint Presentation"
        prs.save(str(path))

        doc = PptxParser().parse(path, "bullets.pptx")
        assert doc.title == "Introduction to Graphs"

    def test_corrupt_file_raises_value_error(self, tmp_path: Path) -> None:
        path = tmp_path / "corrupt.pptx"
        path.write_bytes(b"not a real pptx file")

        with pytest.raises(ValueError, match="PowerPoint"):
            PptxParser().parse(path, "corrupt.pptx")


# ---------------------------------------------------------------------------
# DOCX fixtures
# ---------------------------------------------------------------------------


def _make_docx_nested_headings(tmp_path: Path) -> Path:
    document = Document()
    document.add_paragraph("Some preamble text before any heading.")

    document.add_heading("Week 3", level=1)
    document.add_paragraph("Intro to week 3.")

    document.add_heading("Smoothing", level=2)
    document.add_paragraph("What smoothing is.")

    document.add_heading("Add-one", level=3)
    document.add_paragraph("Add-one smoothing details.")

    document.add_heading("Back to week level", level=1)
    document.add_paragraph("Sibling heading content.")

    path = tmp_path / "nested.docx"
    document.save(str(path))
    return path


def _make_docx_lists_and_table(tmp_path: Path) -> Path:
    document = Document()
    document.add_heading("Shopping", level=1)
    document.add_paragraph("Bananas", style="List Bullet")
    document.add_paragraph("Apples", style="List Bullet")
    document.add_paragraph("Regular paragraph after list.")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Col A"
    table.cell(0, 1).text = "Col B"
    table.cell(1, 0).text = "1"
    table.cell(1, 1).text = "2"

    document.add_paragraph("Text after the table.")

    path = tmp_path / "lists_table.docx"
    document.save(str(path))
    return path


def _make_docx_title_style(tmp_path: Path) -> Path:
    document = Document()
    document.add_paragraph("COMP4650 Notes", style="Title")
    document.add_paragraph("Some content under the title.")
    document.add_heading("Section A", level=1)
    document.add_paragraph("Section A content.")

    path = tmp_path / "titled.docx"
    document.save(str(path))
    return path


def _make_docx_no_title_no_heading1(tmp_path: Path) -> Path:
    document = Document()
    document.add_heading("Subsection only", level=2)
    document.add_paragraph("Content.")

    path = tmp_path / "no_h1.docx"
    document.save(str(path))
    return path


class TestDocxParser:
    def test_nested_headings_produce_heading_path(self, tmp_path: Path) -> None:
        path = _make_docx_nested_headings(tmp_path)
        doc = DocxParser().parse(path, "nested.docx")

        heading_paths = [s.location.heading_path for s in doc.sections]
        assert [] in heading_paths  # preamble section
        assert ["Week 3"] in heading_paths
        assert ["Week 3", "Smoothing"] in heading_paths
        assert ["Week 3", "Smoothing", "Add-one"] in heading_paths
        assert ["Back to week level"] in heading_paths

        add_one_section = next(
            s for s in doc.sections if s.location.heading_path == ["Week 3", "Smoothing", "Add-one"]
        )
        assert add_one_section.title == "Add-one"
        assert "Add-one smoothing details." in add_one_section.text

    def test_preamble_before_first_heading(self, tmp_path: Path) -> None:
        path = _make_docx_nested_headings(tmp_path)
        doc = DocxParser().parse(path, "nested.docx")

        preamble = doc.sections[0]
        assert preamble.location.heading_path == []
        assert preamble.title is None
        assert "Some preamble text before any heading." in preamble.text

    def test_lists_prefixed_and_table_interleaved_in_order(self, tmp_path: Path) -> None:
        path = _make_docx_lists_and_table(tmp_path)
        doc = DocxParser().parse(path, "lists_table.docx")

        assert len(doc.sections) == 1
        text = doc.sections[0].text
        assert "- Bananas" in text
        assert "- Apples" in text
        assert "Regular paragraph after list." in text
        assert "Col A | Col B" in text
        assert "1 | 2" in text

        # Order preserved: list, then plain paragraph, then table, then trailing text.
        idx_list = text.index("- Bananas")
        idx_plain = text.index("Regular paragraph after list.")
        idx_table = text.index("Col A | Col B")
        idx_after = text.index("Text after the table.")
        assert idx_list < idx_plain < idx_table < idx_after

    def test_title_fallback_chain_title_style(self, tmp_path: Path) -> None:
        path = _make_docx_title_style(tmp_path)
        doc = DocxParser().parse(path, "titled.docx")
        assert doc.title == "COMP4650 Notes"

    def test_title_fallback_chain_heading1(self, tmp_path: Path) -> None:
        path = _make_docx_nested_headings(tmp_path)
        doc = DocxParser().parse(path, "nested.docx")
        assert doc.title == "Week 3"

    def test_title_fallback_chain_stem(self, tmp_path: Path) -> None:
        path = _make_docx_no_title_no_heading1(tmp_path)
        doc = DocxParser().parse(path, "no_h1.docx")
        assert doc.title == "no_h1"

    def test_core_properties_title_used_when_present(self, tmp_path: Path) -> None:
        path = _make_docx_nested_headings(tmp_path)
        document = Document(str(path))
        document.core_properties.title = "COMP4650 Week 3 Notes"
        document.save(str(path))

        doc = DocxParser().parse(path, "nested.docx")
        assert doc.title == "COMP4650 Week 3 Notes"

    def test_corrupt_file_raises_value_error(self, tmp_path: Path) -> None:
        path = tmp_path / "corrupt.docx"
        path.write_bytes(b"not a real docx file")

        with pytest.raises(ValueError, match="Word"):
            DocxParser().parse(path, "corrupt.docx")


# ---------------------------------------------------------------------------
# Registry integration
# ---------------------------------------------------------------------------


class TestRegistry:
    def test_registry_lookup_case_insensitive_after_import(self, tmp_path: Path) -> None:
        # Import only the module this builder owns; the other built-in parser
        # modules (pdf, transcripts) may not exist yet in parallel tracks.
        import notecast.ingest.parsers.office  # noqa: F401

        pptx_parser = get_parser(Path("SLIDES.PPTX"))
        assert isinstance(pptx_parser, PptxParser)

        docx_parser = get_parser(Path("notes.DOCX"))
        assert isinstance(docx_parser, DocxParser)
